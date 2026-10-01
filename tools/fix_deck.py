"""Post-conversion fix-ups for the C699 deck: strip the WSQ layer convert_deck.py left behind."""
import copy, re
from pptx import Presentation

PATH = "courseware/QuickBooks Accounting Masterclass (C699)-v1.0.pptx"
SRC = "source-wsq/courseware/Advanced Transactional Accounting with QuickBooks Online-v11.1.pptx"
p = Presentation(PATH)
src = Presentation(SRC)

def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} not found")

def set_para(para, line):
    runs = [r for r in para.runs if r.text] or para.runs
    runs[0].text = line
    for r in para.runs:
        if r is not runs[0] and r._r is not runs[0]._r:
            r.text = ""

def set_paras(slide, sid, lines):
    paras = shape(slide, sid).text_frame.paragraphs
    assert len(paras) == len(lines), (sid, len(paras), lines)
    for para, line in zip(paras, lines):
        set_para(para, line)

def find(title):
    hits = [s for s in p.slides if any(sh.has_text_frame and sh.text_frame.text == title for sh in s.shapes)]
    assert len(hits) == 1, (title, len(hits))
    return hits[0]

# ---- cover
cover = p.slides[0]
set_paras(cover, 8, ["NON-WSQ COURSE SLIDES"])
code_box = shape(cover, 11)   # convert_deck flattens its 3 styled paragraphs: restore from source, drop the TSC line
code_box._element.replace(code_box._element.txBody, copy.deepcopy(shape(src.slides[0], 11)._element.txBody))
tf = code_box.text_frame
set_para(tf.paragraphs[0], "Course Code: C699")
tf.paragraphs[1]._p.getparent().remove(tf.paragraphs[1]._p)

# ---- Lesson Plan slides: 9:30am-5:30pm, 7.5 h, one 30-min lunch (blocks match the LP schedule tables)
d1 = find("Lesson Plan · Day 1"); d2 = find("Lesson Plan · Day 2")
for s in (d1, d2):
    set_paras(s, 4, ["9:30 AM–5:30 PM · 7.5 HOURS"])
set_paras(d1, 11, ["09:30–11:00", "Admin, prerequisite bridge, control-ready company file, double-entry decision model"])
set_paras(d1, 16, ["11:00–13:15", "Advanced order-to-cash, procure-to-pay, Activities 1–3"])
set_paras(d1, 21, ["13:45–16:00", "Recurring transactions, inventory, projects, billable time and costs"])
set_paras(d1, 26, ["16:00–17:30", "Bank rules, reconciliation, month-end adjustments, close and audit controls"])
set_paras(d2, 11, ["09:30–11:30", "Statement architecture, subledger ties, report dimensions and custom reports"])
set_paras(d2, 16, ["11:30–13:10", "Ageing, profitability, cash flow, GST, management reporting"])
set_paras(d2, 21, ["13:40–16:15", "Activities 10–11; Singapore framework, accrual, revenue, inventory, PPE and FX"])
set_paras(d2, 26, ["16:15–17:30", "Estimates, qualitative characteristics, Activity 12 capstone and course review"])

# ---- inline WSQ leftovers
FIXED = {
    "Respect attendance, break, assessment, and submission times.": "Respect start, break, and submission times.",
    "2 days/16 hours; LO1-LO3; published topic and assessment scope": "2 days/15 hours; LO1-LO3; published topics",
    "A4 DECISION LENS": "DECISION LENS",
    "IMPACT ASSESSMENT": "IMPACT REVIEW",
    "Standards impact assessment": "Standards impact review",
    "  All K3/A4 impacts are addressed": "  All standards impacts are addressed",
}
SUMMARY = {"A1": "Transactions", "A2": "Statements", "A3": "Analysis", "A4": "Standards"}
EYEBROW = re.compile(r"(?:\s*·\s*[KA]\d+(?:,\s*[KA]\d+)*)+$")   # "TOPIC 1 · K1 · A1", "ACTIVITY 10 · ADVANCED · K2, A2, A3"
summary = find("What You Achieved")
for s in p.slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.text in FIXED:
                    r.text = FIXED[r.text]
            t = "".join(r.text for r in para.runs)
            if s is summary and t in SUMMARY:
                set_para(para, SUMMARY[t])
            elif re.match(r"^(TOPIC|ACTIVITY) ", t):
                new = EYEBROW.sub("", t)
                if new != t:
                    set_para(para, new)

# ---- Course Feedback: restore the TRAQOM slide's text bodies from the source, then rewrite in place
fb = p.slides[-2]
srcfb = src.slides[136]
for sid in (4, 5, 11, 16, 21, 26):
    dst, s0 = shape(fb, sid), shape(srcfb, sid)
    dst._element.replace(dst._element.txBody, copy.deepcopy(s0._element.txBody))
set_paras(fb, 4, ["COURSE FEEDBACK"])
set_paras(fb, 5, ["Course Feedback & Certificate"])
set_paras(fb, 11, ["Feedback", "Complete the course feedback survey before you leave."])
set_paras(fb, 16, ["Improve", "Tell us what worked and what to improve — it shapes the next run."])
set_paras(fb, 21, ["Support", "Ask the trainer if you cannot open the survey or the course files."])
set_paras(fb, 26, ["Certificate", "Your certificate of completion will be emailed to you after the course."])

# ---- hyperlink relationships still pointing at the WSQ repo
for s in p.slides:
    for rel in s.part.rels.values():
        if rel.is_external and "TGS-2023018989" in rel._target:
            rel._target = rel._target.replace(
                "TGS-2023018989-Advanced-Transactional-Accounting-with-Quickbooks-Online",
                "C699-QuickBooks-Accounting-Masterclass")
p.core_properties.title = "QuickBooks Accounting Masterclass"
p.core_properties.subject = "C699"
p.save(PATH)
print("saved")
