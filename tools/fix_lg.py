"""C699 Learner Guide: explicit rewrite of the WSQ LG (TGS-2023018989) as the non-WSQ twin."""
import re, shutil, docx
from docx.shared import Emu
from docx.oxml.ns import qn
from docxutil import *

SRC = "source-wsq/courseware/LG-Advanced Transactional Accounting with QuickBooks Online.docx"
OUT = "courseware/LG-QuickBooks Accounting Masterclass (C699).docx"
shutil.copy(SRC, OUT)
d = docx.Document(OUT)
retitle_all(d)

# Cover
set_text(para_by_text(d, "WSQ Course Code:"), "Course Code: C699")
delete(para_by_text(d, "TSC: Transactional Accounting"))
set_text(para_by_text(d, "Version 11.1"), "Version 1.0 · 1 October 2026")

t = d.tables
# Version control record -> single v1.0 row
for r in list(t[0].rows)[2:]:
    t[0]._tbl.remove(r._tr)
set_cell(t[0].rows[1].cells[0], "1.0")
set_cell(t[0].rows[1].cells[1], "1 October 2026")
set_cell(t[0].rows[1].cells[2], "Initial release of the QuickBooks Accounting Masterclass learner guide with twelve "
         "progressive QuickBooks Online activities.")
set_cell(t[0].rows[1].cells[3], "Dr Alfred Ang")

# Course information
set_cell(t[1].rows[1].cells[1], "C699")
set_cell(t[1].rows[3].cells[1], "2 days · 15 hours (9:30 am – 5:30 pm)")
t[1]._tbl.remove(t[1].rows[2]._tr)                  # TSC row

# K/A criteria -> removed
set_text(para_by_text(d, "A4 Decision Lens"), "Decision Lens")
set_text(para_by_text(d, "Impact Assessment"), "Impact Review")
for p in d.paragraphs:
    if re.match(r"Level: .+ · Criteria: ", p.text):
        set_text(p, p.text.split(" · Criteria:")[0])

# Detailed Activities overview table: drop the Criteria column, give its width to the Activity column
acts = t[5]
assert acts.rows[0].cells[2].text == "Criteria"
grid = acts._tbl.tblGrid
cols = grid.findall(qn("w:gridCol"))
w = [int(c.get(qn("w:w"))) for c in cols]
for r in acts.rows:
    r._tr.remove(r.cells[2]._tc)
grid.remove(cols[2])
cols[0].set(qn("w:w"), str(w[0] + w[2]))
for r in acts.rows:
    r.cells[0].width = Emu((w[0] + w[2]) * 635)

# Capstone acceptance row
for tb in t:
    for c in tb._cells:
        if c.text == "All K3/A4 impacts are addressed":
            set_cell(c, "All standards impacts are addressed")

# Assessment Preparation section -> removed
assert delete_range(d, "Assessment Preparation", "Source Register") == 5

# Source register: course page row -> the non-WSQ product page
for r in t[-1].rows:
    if r.cells[0].text == "Course page":
        set_cell(r.cells[1], "2 days/15 hours; LO1-LO3; published topics")
        for wt in r.cells[2]._tc.iter(qn("w:t")):     # URL text sits inside <w:hyperlink>, not cell.paragraphs[].runs
            wt.text = wt.text.replace("wsq-advanced-transactional-accounting-with-quickbooks-online",
                                      "quickbooks-accounting-masterclass")

# Hyperlink targets live in the rels, not the text
for rel in d.part.rels.values():
    if rel.is_external and "wsq-advanced-transactional-accounting-with-quickbooks-online" in rel._target:
        rel._target = "https://www.tertiarycourses.com.sg/quickbooks-accounting-masterclass.html"

# Footer
for s in d.sections:
    for part in (s.footer, s.header, s.first_page_footer, s.even_page_footer):
        for p in part.paragraphs:
            for r in p.runs:
                if "TGS-2023018989" in r.text:
                    r.text = r.text.replace("TGS-2023018989", "C699")
d.core_properties.title = "LG-QuickBooks Accounting Masterclass (C699)"
d.save(OUT)
print("saved", OUT)
