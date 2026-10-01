"""C699 Lesson Plan: explicit rewrite of the WSQ LP (TGS-2023018989) as the non-WSQ twin."""
import re, shutil, docx
from docxutil import *

SRC = "source-wsq/courseware/LP-Advanced Transactional Accounting with QuickBooks Online.docx"
OUT = "courseware/LP-QuickBooks Accounting Masterclass (C699).docx"
shutil.copy(SRC, OUT)
d = docx.Document(OUT)
retitle_all(d)

# Kept-slide map: source slides 2, 7, 8, 12, 13, 14, 133, 135, 136 (1-based) were dropped.
DROPPED = [2, 7, 8, 12, 13, 14, 133, 135, 136]
def m(n):
    n = int(n)
    assert n not in DROPPED, n
    return str(n - sum(1 for x in DROPPED if x < n))
def remap(text):
    return re.sub(r"\b\d+", lambda mo: m(mo.group(0)), text)

# Cover
set_text(para_by_text(d, "WSQ Course Code:"), "Course Code: C699")
delete(para_by_text(d, "TSC: Transactional Accounting"))
set_text(para_by_text(d, "Version 11.1"), "Version 1.0 · 1 October 2026")

# Assessment Mapping section (heading + table)
t0, t1, t2, t3, t4, t5 = d.tables
delete(para_by_text(d, "Assessment Mapping"))
t2._tbl.getparent().remove(t2._tbl)

# Version control record -> single v1.0 row
for r in list(t0.rows)[2:]:
    t0._tbl.remove(r._tr)
set_cell(t0.rows[1].cells[0], "1.0")
set_cell(t0.rows[1].cells[1], "1 October 2026")
set_cell(t0.rows[1].cells[2], "Initial release of the QuickBooks Accounting Masterclass lesson plan: two-day schedule "
         "(9:30am–5:30pm, 15 instructional hours) with twelve progressive QuickBooks Online activities.")
set_cell(t0.rows[1].cells[3], "Dr Alfred Ang")

# Course information: code, drop TSC + Assessment rows, duration, timing
set_cell(t1.rows[1].cells[1], "C699")
set_cell(t1.rows[3].cells[1], "2 days · 7.5 instructional hours per day · 15 hours total")
set_cell(t1.rows[4].cells[1], "9:30 am – 5:30 pm · 30-minute lunch")
for r in (t1.rows[5], t1.rows[2]):          # Assessment, TSC
    t1._tbl.remove(r._tr)

s3 = [[c.text for c in r.cells] for r in t3.rows]
s4 = [[c.text for c in r.cells] for r in t4.rows]
# Day 1: teaching rows 460 min rescaled onto 450 (the two 80-min blocks give 5 min each);
# 10-min teas and the 60-min lunch become a single 30-min lunch.
rebuild_rows(t3, [
    (1, ["09:30–10:00", "30", "Admin, learning outcomes and course orientation", "1–10"]),
    (2, ["10:00–11:00", "60", s3[2][2], remap(s3[2][3])]),
    (4, ["11:00–12:15", "75", s3[4][2], s3[4][3]]),
    (5, ["12:15–13:15", "60", s3[5][2], remap(s3[5][3])]),
    (6, ["13:15–13:45", "30", "Lunch", ""]),
    (7, ["13:45–15:00", "75", s3[7][2], s3[7][3]]),
    (9, ["15:00–16:00", "60", s3[9][2], remap(s3[9][3])]),
    (10, ["16:00–16:55", "55", s3[10][2], s3[10][3]]),
    (11, ["16:55–17:30", "35", "Close/audit controls; Activities 6–9 briefing; Day 1 recap", remap(s3[11][3])]),
])
# Day 2: the 120-min assessment block and teas removed; the 340 min of teaching rows rescaled
# proportionally onto 450 (x1.32, 5-min granularity).
rebuild_rows(t4, [
    (1, ["09:30–09:50", "20", "Day 1 recap", "Recap"]),
    (2, ["09:50–11:30", "100", s4[2][2], remap(s4[2][3])]),
    (4, ["11:30–13:10", "100", s4[4][2], s4[4][3]]),
    (6, ["13:10–13:40", "30", "Lunch", ""]),
    (5, ["13:40–15:00", "80", s4[5][2], remap(s4[5][3])]),
    (7, ["15:00–16:15", "75", s4[7][2], remap(s4[7][3])]),
    (9, ["16:15–17:10", "55", s4[9][2], remap(s4[9][3])]),
    (10, ["17:10–17:30", "20", "Consolidation, course feedback and Q&A", "126–129"]),
])
for p in [p for p in d.paragraphs if p.text.startswith("Scheduled non-lunch time:")]:
    set_text(p, "Instructional time: 450 minutes (7.5 hours), 9:30am – 5:30pm with a single 30-minute lunch.")

# Activity traceability: drop the Criteria (K/A) column, remap slide ranges
from docx.shared import Emu
grid = t5._tbl.tblGrid
widths = [int(g.get(docx.oxml.ns.qn("w:w"))) for g in grid.findall(docx.oxml.ns.qn("w:gridCol"))]
for r in t5.rows:
    r._tr.remove(r.cells[2]._tc)
grid.remove(grid.findall(docx.oxml.ns.qn("w:gridCol"))[2])
freed = widths[2]
cols = grid.findall(docx.oxml.ns.qn("w:gridCol"))
keep = [w for i, w in enumerate(widths) if i != 2]
new = [w + freed * w // sum(keep) for w in keep]
for g, w in zip(cols, new):
    g.set(docx.oxml.ns.qn("w:w"), str(w))
for r in t5.rows:
    for c, w in zip(r.cells, new):
        c.width = Emu(w * 635)
for r in t5.rows[1:]:
    set_cell(r.cells[2], remap(r.cells[2].text))
set_cell(t5.rows[12].cells[3], "Standards impact review")
heading = para_by_text(d, "Activity and Criteria Traceability")
set_text(heading, "Activity Traceability")

# Footer
for s in d.sections:
    for part in (s.footer, s.header, s.first_page_footer, s.even_page_footer):
        for p in part.paragraphs:
            for r in p.runs:
                if "TGS-2023018989" in r.text:
                    r.text = r.text.replace("TGS-2023018989", "C699")
d.core_properties.title = "LP-QuickBooks Accounting Masterclass (C699)"
d.save(OUT)
print("saved", OUT)
